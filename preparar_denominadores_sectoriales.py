"""Descarga bases oficiales y conserva los cruces aceptados y rechazados.

No imputa municipios. Los registros administrativos no se homologan a los daños
por similitud de nombres. Generar el tablero usa el JSON guardado, sin red.
"""
import hashlib
import io
import json
import math
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import urlopen

ROOT = Path(__file__).resolve().parent
PAGE = 'https://www.dane.gov.co/index.php/estadisticas-por-tema/demografia-y-poblacion/proyecciones-de-viviendas-y-hogares'
VIV = 'https://www.dane.gov.co/files/operaciones/PPVH/anexo-proyecciones-viviendas-dptal-mpal-2018-2042.xlsx'
HOG = 'https://www.dane.gov.co/files/operaciones/PPVH/anexo-proyecciones-hogares-dptal-mpal-2018-2042.xlsx'
REPS = 'https://www.datos.gov.co/resource/c36g-9fc2.json'
YEAR = 2026


def download(url):
    raw = urlopen(url, timeout=60).read()
    return raw, hashlib.sha256(raw).hexdigest()


def extract_projection(raw, sheet_name, kind, source, year=YEAR):
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = book[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[9]
    col = next(i for i, v in enumerate(header) if str(v) == str(year))
    records, seen = [], set()
    for line, values in enumerate(rows, 1):
        if len(values) <= col or str(values[4]).strip().casefold() != 'total':
            continue
        code = str(values[2]).zfill(5)
        if len(code) != 5 or not code.isdigit():
            continue  # Exclude footnotes, never guess municipality identifiers.
        if code in seen:
            raise ValueError(f'Denominador duplicado: {kind} {code}')
        seen.add(code)
        value = values[col]
        if not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
            continue
        records.append(dict(code=code, year=year, kind=kind, value=value,
                            unit='Viviendas' if kind == 'viviendas' else 'Hogares',
                            source=source, reference_date=f'{year}-06-30',
                            locator=f'{sheet_name}!{openpyxl.utils.get_column_letter(col+1)}{line}',
                            area='Total', status='verified'))
    book.close()
    if len(records) < 1100:
        raise ValueError(f'Cobertura inesperada de {kind}: {len(records)}')
    return records


def main():
    sources, records = {}, []
    for id, url, sheet, kind in [
        ('dane_viviendas', VIV, 'Proye total viviendas mpio', 'viviendas'),
        ('dane_hogares', HOG, 'Proyecciones Hogares mpio', 'hogares'),
    ]:
        raw, sha = download(url)
        records += extract_projection(raw, sheet, kind, id)
        sources[id] = dict(label=f'DANE · Proyección municipal de {kind} 2026', url=url,
                           page=PAGE, published='2025-12-24', sha256=sha,
                           reference='2026-06-30', method='Proyección oficial; área Total')

    population = json.loads((ROOT / 'data/poblacion_relativa.json').read_text(encoding='utf-8'))
    # Re-download and compare the original file, so a stale snapshot is not silently certified.
    _, population_sha = download(population['url'])
    if population_sha != population['sha256']:
        raise ValueError('DANE cambió el archivo de población: revisar y regenerar antes de continuar')
    sources['dane_poblacion'] = dict(label=population['source'], url=population['url'], page=population['page'],
                                   published=population['publication'], sha256=population_sha,
                                   reference='2026-06-30', method='Proyección oficial; área Total')
    for r in population['rows']:
        if r['year'] == YEAR:
            records.append(dict(code=r['code'], year=r['year'], kind='poblacion', value=r['population'],
                                unit='Habitantes', source='dane_poblacion', reference_date='2026-06-30',
                                locator=r['locator'], area='Total', status='verified'))

    # Count site identifiers by LOCATION OF THE SITE, never by provider headquarters.
    reps_filter = "claseprestador='Instituciones Prestadoras de Servicios de Salud - IPS' OR municipiosede in ('27050','27660')"
    query = REPS + '?' + urlencode({'$select':'codigohabilitacionsede,nombresede,municipiosede,municipiosededesc,claseprestador,fecha_corte_reps',
                                    '$where':reps_filter, '$order':'codigohabilitacionsede', '$limit':50000})
    raw, sha = download(query)
    reps = json.loads(raw)
    count = json.loads(download(REPS + '?' + urlencode({'$select':'count(*) as n', '$where':reps_filter}))[0])[0]
    if len(reps) != int(count['n']):
        raise ValueError('Descarga REPS incompleta: no producir totales parciales')
    cutoffs = sorted({r.get('fecha_corte_reps', '') for r in reps})
    if cutoffs != ['Fecha corte REPS: Mar 12 2026  3:11PM']:
        raise ValueError(f'Cambió el corte REPS; revisar su fecha y alcance: {cutoffs}')
    groups = defaultdict(set)
    identities = defaultdict(set)
    for r in reps:
        code, site = r.get('municipiosede', ''), r.get('codigohabilitacionsede', '')
        if site:
            identities[site].add((code, r.get('claseprestador')))
        if len(code) == 5 and code.isdigit() and site and r.get('claseprestador') == 'Instituciones Prestadoras de Servicios de Salud - IPS':
            groups[code].add(site)
    conflicts = {s for s, keys in identities.items() if len(keys) > 1}
    counts = [dict(code=code, value=len(sites), kind='sedes_ips', unit='Sedes IPS', source='reps', reference_date='2026-03-12',
                   status='candidate_only', reason='No se homologaron los puntos/centros afectados a sedes IPS REPS')
              for code, sites in sorted(groups.items()) if not sites.intersection(conflicts)]
    sources['reps'] = dict(label='Minsalud · REPS, sedes IPS', url=query,
                           page='https://www.datos.gov.co/d/c36g-9fc2', reference='2026-03-12',
                           sha256=sha, rows=len(reps), unique_sites=len(identities),
                           method='Sedes IPS distintas por codigohabilitacionsede y municipiosede; no implica edificios únicos ni operación tras el evento')

    # Keep publisher and ACTUAL DATA years; metadata updates are not new observations.
    source_audits = []
    for id in ['x5ay-984n', 'j9sd-zau5', '3ncw-3qwq']:
        metadata_url = f'https://www.datos.gov.co/api/views/{id}.json'
        raw_meta, meta_sha = download(metadata_url)
        meta = json.loads(raw_meta)
        entry = dict(id=id, url=metadata_url, sha256=meta_sha, name=meta['name'],
                     publisher=meta.get('attribution'), data_updated=meta.get('rowsUpdatedAt'),
                     description=meta.get('description'), status='not_accepted')
        if id != '3ncw-3qwq':
            url = f'https://www.datos.gov.co/resource/{id}.json?' + urlencode({'$select':'a_o,count(*) as n', '$group':'a_o', '$order':'a_o'})
            raw_years, year_sha = download(url)
            entry.update(years=json.loads(raw_years), years_url=url, years_sha256=year_sha)
        source_audits.append(entry)

    payload = dict(version='sectorial-1', year=YEAR, event_date='2026-08-10',
                   checked=datetime.now(timezone.utc).isoformat(timespec='seconds'),
                   sources=sources, rows=sorted(records, key=lambda r:(r['kind'],r['code'])),
                   candidates=counts,
                   audit=dict(reps_targets=[r for r in reps if r.get('municipiosede') in ('27050','27660')],
                              reps_conflicting_site_ids=sorted(conflicts), education=source_audits))
    target=ROOT / 'data/denominadores_sectoriales.json'
    target.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':'), allow_nan=False), encoding='utf-8')
    print(f'{len(records)} bases DANE verificadas; {len(counts)} inventarios IPS candidatos, sin habilitar ratios de salud')
    for code in ['27050','27660']:
        print(code, [(r['kind'],r['value']) for r in records if r['code']==code], [r for r in counts if r['code']==code])
    print('Años educativos', [(r['id'],r.get('years')) for r in source_audits])


if __name__ == '__main__':
    main()
