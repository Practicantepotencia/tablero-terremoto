"""Contexto de fuentes: agregados seguros, sin sumar universos incompatibles."""
import argparse
from collections import defaultdict
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
try:
    from .preparar_educacion_men import ROOT, norm, code, count
except ImportError:
    from preparar_educacion_men import ROOT, norm, code, count

def prepare(folder):
    folder=Path(folder)
    reference=json.loads((ROOT/'data/poblacion_relativa.json').read_text(encoding='utf-8'))['rows']
    names={(norm(r['d']),norm(r['m'])):r['code'] for r in reference}
    names.update({('risaralda','dos quebradas'):'66170',('valle del cauca','cali'):'76001',
                  ('choco','bajo baudo'):'27077',('choco','alto baudo'):'27025'})
    men=load_workbook(folder/'men_sedes_escolares_afectadas_20260921.xlsx',read_only=True,data_only=True)
    men_rows={code(r[9],12):r for r in list(men['Sheet1'].values)[1:] if r[9] is not None};men.close()
    def open_source(file,sheet,label,unit,note):
        p=folder/file;w=load_workbook(p,read_only=True,data_only=True)
        return w,{'file':file,'sheet':sheet,'label':label,'unit':unit,'source_date':None,
                  'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'note':note,'municipalities':[]}
    sources=[]
    w,s=open_source('icbf_base_afectaciones_aliados.xlsx','PRIMERA INFANCIA','ICBF · primera infancia','UDS',
        'UDS y beneficiarios reportados, no sedes MEN. Beneficiarios desconocidos no se completan. Otras hojas administrativas, SRD y nutrición no se suman a este universo.')
    groups=defaultdict(list);unresolved=0;seen=set()
    for r in list(w[s['sheet']].values)[1:]:
        if r[3] is None:continue
        key=str(r[3]);municipal=names.get((norm(r[0]),norm(r[9])))
        if key in seen:raise ValueError('UDS duplicada: conciliar')
        seen.add(key)
        if not municipal:unresolved+=1;continue
        groups[municipal].append(r)
    for c,rs in sorted(groups.items()):
        values=[count(r[6]) for r in rs]
        s['municipalities'].append({'code':c,'records':len(rs),'units':len(rs),'beneficiaries_reported':sum(v for v in values if v is not None),
            'text_integer_beneficiaries':sum(isinstance(r[6],str) and count(r[6]) is not None for r in rs),'missing_beneficiaries':sum(v is None for v in values),'missing_damage':sum(not norm(r[8]) for r in rs),'severe_units':sum(norm(r[8])=='severa' for r in rs)})
    s['unresolved_records']=unresolved;s['records']=len(seen);w.close();sources.append(s)
    w,s=open_source('sedes_afectadas_filtradas_pei_dqdas.xlsx','sedes_afectadas_Pei y Dqdas','Pereira y Dosquebradas · afectadas','Sedes',
        'Sedes ya incluidas en MEN. Sus categorías y matrículas pueden discrepar; no se añaden ni sustituyen automáticamente. La hoja de proyecto es otro listado.')
    groups=defaultdict(list)
    for r in list(w[s['sheet']].values)[1:]:
        if r[3] is not None:groups[names.get((norm(r[0]),norm(r[1])))].append(r)
    for c,rs in sorted(groups.items(),key=str):
        ids=[code(r[3],12) for r in rs]
        if len(ids)!=len(set(ids)):raise ValueError('Sede municipal duplicada: conciliar')
        matched=[(r,men_rows[code(r[3],12)]) for r in rs if code(r[3],12) in men_rows]
        s['municipalities'].append({'code':c,'records':len(rs),'units':len(ids),'men_matches':len(matched),
          'enrollment_reported':sum(count(r[7]) or 0 for r in rs),'men_enrollment_same_sites':sum(count(m[38]) or 0 for r,m in matched),
          'damage_disagreements':sum(norm(r[16])!=norm(m[41]) for r,m in matched),
          'enrollment_disagreements':sum(count(r[7])!=count(m[38]) for r,m in matched)})
    project=[r for r in list(w['sedes_proyecto_Pei y Dqdas'].values)[1:] if r[3] is not None]
    s['project_records']=len(project);s['project_unique_codes']=len({code(r[3],12) for r in project});w.close();sources.append(s)
    w,s=open_source('pereira_fuente_unica_v2_dane.xlsx','Matriz_maestra','Pereira · fuente única','Instituciones',
        'El DANE identifica instituciones, no sedes. La matrícula institucional repetida no se suma por fila; requiere concordancia institución-sede para cruzar MEN.')
    rs=[r for r in list(w[s['sheet']].values)[1:] if any(x is not None for x in r)]
    codes={code(r[3],12) for r in rs if r[3] is not None}
    s['municipalities']=[{'code':'66001','records':len(rs),'units':len(codes),'missing_code':sum(r[3] is None for r in rs)}]
    w.close();sources.append(s)
    w,s=open_source('fundacion_plan_sedes.xlsx','Hoja1','Fundación PLAN','Sedes',
        'Intervenciones reportadas en sedes ya presentes en MEN. No acreditan ejecución, cobertura efectiva ni beneficiarios atendidos.')
    groups=defaultdict(list)
    for r in list(w[s['sheet']].values)[3:]:
        if r[5] is not None:groups[names.get((norm(r[1]),norm(r[2])))].append(r)
    for c,rs in sorted(groups.items(),key=str):
        ids={code(r[5],12) for r in rs}
        s['municipalities'].append({'code':c,'records':len(rs),'units':len(ids),'men_matches':sum(i in men_rows for i in ids)})
    w.close();sources.append(s)
    return {'version':1,'reference_commit':'25c13842159a57d8fbeffb8a5ac8b51d96e14a6b',
            'purpose':'Contexto por unidad propia; no alimenta el índice ni estima cobertura exhaustiva.',
            'sources':sources}

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('folder',type=Path)
    ap.add_argument('--out',type=Path,default=ROOT/'data/fuentes_educacion_contexto.json');a=ap.parse_args()
    out=prepare(a.folder);a.out.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps([{'source':s['label'],'municipalities':len(s['municipalities']),'records':sum(r['records'] for r in s['municipalities']),'units':sum(r['units'] for r in s['municipalities']),'unresolved':s.get('unresolved_records',sum(r['code'] is None for r in s['municipalities']))} for s in out['sources']],ensure_ascii=False))
