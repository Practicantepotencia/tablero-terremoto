"""Agregados municipales MEN. No exporta filas de sedes ni datos de personas."""
import argparse
from collections import Counter, defaultdict
import hashlib
import json
import math
from pathlib import Path
import unicodedata

ROOT=Path(__file__).resolve().parents[1]
CRITICAL={'colapso total','colapso parcial','riesgo inminente de colapso'}
KNOWN=CRITICAL|{'afectacion menor','afectacion parcial','sin afectacion'}
FIVE={'17','27','63','66','76'}

def norm(x):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFD',str(x or '').casefold()) if not unicodedata.combining(c)).split())

def code(x,n):
    if isinstance(x,bool):raise ValueError('Código booleano')
    s=str(int(x)) if isinstance(x,(int,float)) and math.isfinite(x) and x==int(x) else str(x or '').strip()
    if not s.isdigit() or len(s)>n:raise ValueError('Código no válido')
    return s.zfill(n)

def count(x):
    # Some source workbooks encode integer counts as text; blanks remain unknown.
    if isinstance(x,str) and x.strip().isdigit():return int(x.strip())
    if isinstance(x,bool) or not isinstance(x,(int,float)) or not math.isfinite(x) or x<0 or x!=int(x):return None
    return int(x)

def aggregate(records):
    groups=defaultdict(list);seen=set()
    for r in records:
        if r['site'] in seen:raise ValueError('Código de sede MEN duplicado; conciliar antes de sumar')
        seen.add(r['site']);groups[r['code']].append(r)
    output=[]
    for municipal,rs in sorted(groups.items()):
        critical=[r for r in rs if r['damage'] in CRITICAL]
        missing=sum(r['enrollment'] is None for r in critical)
        unknown=sum(r['damage'] not in KNOWN for r in rs)
        total_missing=sum(r['enrollment'] is None for r in rs)
        row={'code':municipal,'d':rs[0]['d'],'m':rs[0]['m'],'reported_sites':len(rs),
             'affected_sites':sum(r['damage'] in KNOWN-{'sin afectacion'} for r in rs),
             'institutions':len({r['institution'] for r in rs}),
             'enrollment':None if total_missing else sum(r['enrollment'] for r in rs),
             'critical_sites':len(critical),'critical_enrollment':None if missing or unknown else sum(r['enrollment'] for r in critical),
             'missing_enrollment':total_missing,'missing_critical_enrollment':missing,'unclassified_sites':unknown,
             'urban_sites':sum(r['zone']=='urbana' for r in rs),'rural_sites':sum(r['zone']=='rural' for r in rs)}
        for status in ['yes','no','unknown']:
            found=[r for r in rs if r['service']==status]
            row['service_'+status+'_sites']=len(found)
            row['service_'+status+'_enrollment']=None if any(r['enrollment'] is None for r in found) else sum(r['enrollment'] for r in found)
        row['critical_service_yes_sites']=sum(r['service']=='yes' for r in critical)
        row['critical_service_yes_enrollment']=None if any(r['enrollment'] is None for r in critical if r['service']=='yes') else sum(r['enrollment'] for r in critical if r['service']=='yes')
        output.append(row)
    return output

def prepare(path):
    from openpyxl import load_workbook
    path=Path(path)
    reference=json.loads((ROOT/'data/poblacion_relativa.json').read_text(encoding='utf-8'))
    refs={r['code']:r for r in reference['rows']}
    book=load_workbook(path,read_only=True,data_only=True);sheet=book['Sheet1']
    rows=sheet.iter_rows(values_only=True);headers=next(rows)
    for i,label in [(3,'CODIGO DANE MUNICIPIO'),(5,'CODIGO DANE INSTITUCION EDUCATIVA'),(9,'CODIGO DANE SEDE'),(38,'MATRICULA TOTAL NACIONAL')]:
        if norm(headers[i])!=norm(label):raise ValueError('El esquema MEN cambió: '+label)
    records=[]
    for row in rows:
        if not any(v is not None for v in row):continue
        municipal=code(row[3],5)
        if municipal not in refs or norm(refs[municipal]['d'])!=norm(row[1]):raise ValueError('Municipio MEN no homologado')
        if norm(row[7])!='oficial':raise ValueError('Sector educativo no previsto')
        records.append({'code':municipal,'d':refs[municipal]['d'],'m':refs[municipal]['m'],
            'site':code(row[9],12),'institution':code(row[5],12),'enrollment':count(row[38]),
            'damage':norm(row[41]),'zone':norm(row[11]),
            'service':{'si':'yes','no':'no'}.get(norm(row[49]),'unknown')})
    book.close();municipalities=aggregate(records)
    roster=[{k:r[k] for k in ['code','d','m']} for c,r in sorted(refs.items()) if c[:2] in FIVE]
    totals={k:None if any(r[k] is None for r in municipalities) else sum(r[k] for r in municipalities) for k in ['reported_sites','critical_sites','critical_enrollment','enrollment','service_yes_sites','service_no_sites','service_unknown_sites','critical_service_yes_sites','critical_service_yes_enrollment']}
    totals.update(municipalities=len(municipalities),institutions=len({r['institution'] for r in records}),departments=len({r['d'] for r in records}))
    return {'version':'MEN-20260921.1','enabled':True,'effective_from':'2026-09-21','report_date':'2026-09-21',
            'report_date_basis':'Inferida del nombre del archivo; no es fecha de inspección ni de matrícula.',
            'observed_at':None,'critical_share':0.5,'reference_commit':'25c13842159a57d8fbeffb8a5ac8b51d96e14a6b',
            'source':{'file':path.name,'sheet':'Sheet1','sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                      'url':'https://github.com/Practicantepotencia/tablero-terremoto/tree/educacion-matricula-critica',
                      'locator':'D municipio; F institución; J sede; AM matrícula; AP daño; AX servicio'},
            'critical_damage':sorted(CRITICAL),'municipalities':municipalities,'roster':roster,'totals':totals}

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('input',type=Path)
    ap.add_argument('--out',type=Path,default=ROOT/'data/educacion_men.json');args=ap.parse_args()
    result=prepare(args.input)
    args.out.write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
    print(json.dumps(result['totals'],ensure_ascii=False))
